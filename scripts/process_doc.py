# -*- coding: utf-8 -*-
import docx
import json
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = r'd:\SVTHM2026'
DOC_PATH = os.path.join(ROOT_DIR, 'SERIES THỔ CẨM.docx')

doc = docx.Document(DOC_PATH)

# ==============================================================================
# 1. TRÍCH XUẤT 20 HOA VĂN (Ý NGHĨA HOA VĂN)
# ==============================================================================
motif_names = [
    'Con gà', 'Hình Học cổ', 'Cây Lúa', 'Hoa Mào Gà', 'Cây Nêu Yang',
    'Con cá suối', 'Cây tre', 'Cây cỏ núi rừng', 'Chậu hoa', 'Ghè cổ',
    'Giã Gạo Ban Mai', 'Gánh Buôn', 'Dấu Xưa Di Sản', 'Sinh hoạt buôn làng',
    'Mẹ và Chị', 'Người Làm Rẫy', 'Đám trẻ thơ', 'Đôi lứa',
    'Khát vọng vươn xa', 'Máy bay'
]

ranges = [
    (1, 49, 54),
    (2, 55, 65),
    (3, 66, 70),
    (4, 71, 75),
    (5, 76, 80),
    (6, 81, 86),
    (7, 87, 92),
    (8, 93, 102),
    (9, 103, 108),
    (10, 109, 115),
    (11, 116, 121),
    (12, 122, 128),
    (13, 129, 141),
    (14, 142, 147),
    (15, 148, 156),
    (16, 157, 164),
    (17, 165, 176),
    (18, 177, 182),
    (19, 183, 188),
    (20, 189, 194),
]

motifs_dict = {}

for m_num, start_idx, end_idx in ranges:
    paras = [doc.paragraphs[k].text.strip() for k in range(start_idx, end_idx) if doc.paragraphs[k].text.strip()]
    
    # Overview is usually paras[0]
    overview = paras[0] if paras else ''
    bullets = []
    sections = {}
    current_sec = 'Chi tiết'
    
    for p in paras[1:]:
        if p.startswith(('1.', '2.', '3.', '4.', 'Ý nghĩa', 'Hệ thống')):
            current_sec = p
            if current_sec not in sections:
                sections[current_sec] = []
        else:
            if current_sec not in sections:
                sections[current_sec] = []
            sections[current_sec].append(p)
            bullets.append(p)
            
    motifs_dict[f'motif-jrai-{m_num:03d}'] = {
        'id': f'motif-jrai-{m_num:03d}',
        'num': m_num,
        'name': motif_names[m_num-1],
        'overview': overview,
        'bullets': bullets,
        'sections': sections,
        'full_text': '\n\n'.join(paras)
    }

# ==============================================================================
# 2. TRÍCH XUẤT 20 CÂU CHUYỆN SẢN PHẨM
# ==============================================================================
story_starts = [
    235, 253, 275, 300, 322, 340, 360, 388, 406, 432, 453, 476, 501, 519, 545, 566, 595, 636, 652, 670
]

stories = []
for idx, st in enumerate(story_starts):
    end = story_starts[idx+1] if idx+1 < len(story_starts) else len(doc.paragraphs)
    title = doc.paragraphs[st].text.strip()
    paras = [doc.paragraphs[k].text.strip() for k in range(st+1, end) if doc.paragraphs[k].text.strip()]
    lead = paras[0] if paras else ''
    body_paras = paras[1:] if len(paras) > 1 else paras
    stories.append({
        'index': idx+1,
        'title': title,
        'lead': lead,
        'paragraphs': body_paras,
        'full_text': '\n\n'.join(paras)
    })

output_data = {
    'motifs': motifs_dict,
    'stories': stories
}

with open(os.path.join(ROOT_DIR, 'scripts', 'series_data.json'), 'w', encoding='utf-8') as f:
    json.dump(output_data, f, ensure_ascii=False, indent=2)

print(f"Successfully processed {len(motifs_dict)} motifs and {len(stories)} stories into series_data.json!")

